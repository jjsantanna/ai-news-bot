"""Tests for ingest loaders.

Skips PDF/DOCX/URL tests when their optional deps aren't installed; the
plaintext loader has zero deps and is always tested.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from compliance_extractor.ingest.base import Document, TextBlock
from compliance_extractor.ingest.plaintext_loader import PlaintextLoader
from compliance_extractor.ingest.registry import get_loader, load


@pytest.mark.fast
def test_plaintext_loader_roundtrip(tmp_path: Path) -> None:
    p = tmp_path / "policy.md"
    p.write_text("# Policy\n\nAccess is restricted to authorized personnel.\n")
    loader = PlaintextLoader()
    assert loader.can_handle(p)
    doc = loader.load(p)
    assert isinstance(doc, Document)
    assert doc.doc_id == "policy.md"
    assert "authorized personnel" in doc.full_text
    assert len(doc.blocks) == 1
    block = doc.blocks[0]
    assert isinstance(block, TextBlock)
    assert doc.full_text[block.char_start : block.char_end] == doc.full_text


@pytest.mark.fast
def test_registry_dispatches_plaintext(tmp_path: Path) -> None:
    p = tmp_path / "x.txt"
    p.write_text("hello\n")
    loader = get_loader(p)
    assert loader.name == "plaintext"
    doc = load(p)
    assert doc.full_text == "hello\n"


@pytest.mark.fast
def test_registry_url_dispatch_no_fetch() -> None:
    """We dispatch UrlLoader for http(s) URLs without actually fetching."""
    loader = get_loader("https://example.com/policy.html")
    assert loader.name == "url"


@pytest.mark.fast
def test_registry_unknown_source_raises(tmp_path: Path) -> None:
    p = tmp_path / "weird.bin"
    p.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError, match="No registered loader"):
        get_loader(p)


@pytest.mark.fast
def test_pdf_loader_with_pypdf(tmp_path: Path) -> None:
    pypdf = pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    pdf_path = tmp_path / "tiny.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf_path.open("wb") as f:
        writer.write(f)

    from compliance_extractor.ingest.pdf_loader import PdfLoader

    loader = PdfLoader()
    assert loader.can_handle(pdf_path)
    doc = loader.load(pdf_path)
    assert doc.metadata["modality"] == "pdf"
    assert doc.metadata["n_pages"] == 1
    assert doc.blocks[0].page == 1
    assert pypdf is not None  # silence unused-import lint


@pytest.mark.fast
def test_docx_loader_preserves_section_path(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document as DocxDocument

    docx_path = tmp_path / "policy.docx"
    d = DocxDocument()
    d.add_heading("Security", level=1)
    d.add_heading("Access Control", level=2)
    d.add_paragraph("Multi-factor authentication is required for all administrators.")
    d.save(str(docx_path))

    from compliance_extractor.ingest.docx_loader import DocxLoader

    doc = DocxLoader().load(docx_path)
    quote = "Multi-factor authentication is required for all administrators."
    assert quote in doc.full_text
    sections = {b.section for b in doc.blocks if b.section}
    assert "Security / Access Control" in sections
