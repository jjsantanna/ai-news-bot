"""DOCX loader using python-docx. Preserves heading-aware section paths.

Each paragraph becomes a TextBlock; the `section` field tracks the most recent
heading hierarchy as a slash-joined path (e.g. `Security / Access Control`).
"""
from __future__ import annotations

from pathlib import Path

from compliance_extractor.ingest.base import Document, TextBlock


class DocxLoader:
    name = "docx"

    def can_handle(self, source: str | Path) -> bool:
        p = Path(str(source))
        return p.suffix.lower() == ".docx" and p.exists()

    def load(self, source: str | Path) -> Document:
        from docx import Document as DocxDocument  # python-docx

        path = Path(source)
        doc = DocxDocument(str(path))

        blocks: list[TextBlock] = []
        full_parts: list[str] = []
        offset = 0
        section_stack: list[tuple[int, str]] = []  # (level, title)

        for para in doc.paragraphs:
            text = para.text or ""
            style_name = (para.style.name if para.style else "") or ""
            if style_name.startswith("Heading "):
                try:
                    level = int(style_name.split(" ", 1)[1])
                except (IndexError, ValueError):
                    level = 1
                section_stack = [(lvl, ttl) for (lvl, ttl) in section_stack if lvl < level]
                if text:
                    section_stack.append((level, text))
            section_path = " / ".join(t for _, t in section_stack) or None

            line = text + "\n"
            blocks.append(
                TextBlock(
                    text=line,
                    char_start=offset,
                    char_end=offset + len(line),
                    page=None,
                    section=section_path,
                )
            )
            full_parts.append(line)
            offset += len(line)

        full_text = "".join(full_parts)
        return Document(
            doc_id=path.name,
            full_text=full_text,
            blocks=blocks,
            metadata={"source_path": str(path), "modality": "docx"},
        )
